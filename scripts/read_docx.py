import sys, subprocess
try:
    from docx import Document
except ImportError:
    subprocess.run([sys.executable, '-m', 'pip', 'install', 'python-docx', '--quiet'])
    from docx import Document

path = r"C:\Users\user\Downloads\WATI_API_Integration__1_ (1).docx"
d = Document(path)
for p in d.paragraphs:
    if p.text.strip():
        print(p.text)
for i, t in enumerate(d.tables):
    print(f"\n--- TABLE {i+1} ---")
    for row in t.rows:
        print(" | ".join(c.text for c in row.cells))
